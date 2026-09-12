"""表格解析器（xlsx / csv）。

表格行是天然的问答素材（培养方案的学分表、课时表都在这里），
所以不按单元格平铺，而是渲染成 "列名: 值 ｜ 列名: 值" 的句子，
让每个数据行都能被自然语言问题检索到。
"""
import csv
import io

from .base import ParseError, ParseResult

MAX_ROWS = 5000  # 防止超大表把内存吃光


def _rows_to_blocks(rows, result: ParseResult, title: str) -> None:
    if not rows:
        return
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    result.add(title, heading_path=title, level=1)

    if len(rows) == 1:
        result.add(" | ".join(h for h in header if h), heading_path=title)
        return

    for row in rows[1:MAX_ROWS + 1]:
        cells = [str(c).strip() if c is not None else "" for c in row]
        pairs = [
            f"{header[i]}: {cells[i]}"
            for i in range(min(len(header), len(cells)))
            if header[i] and cells[i]
        ]
        # 表头缺失时退化成竖线拼接，至少保留内容
        line = " ｜ ".join(pairs) if pairs else " | ".join(c for c in cells if c)
        if line:
            result.add(line, heading_path=title)


def _parse_xlsx(path: str, result: ParseResult) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(values_only=True):
                if any(c is not None and str(c).strip() for c in row):
                    rows.append(list(row))
                if len(rows) > MAX_ROWS:
                    break
            _rows_to_blocks(rows, result, ws.title)
    finally:
        wb.close()


def _parse_csv(path: str, result: ParseResult) -> None:
    raw = None
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "latin-1"):
        try:
            with open(path, "r", encoding=encoding, newline="") as fh:
                raw = fh.read()
            break
        except UnicodeDecodeError:
            continue
    if raw is None:
        raise ParseError("CSV 编码无法识别")

    reader = csv.reader(io.StringIO(raw))
    rows = [r for r in reader if any(str(c).strip() for c in r)][: MAX_ROWS + 1]
    _rows_to_blocks(rows, result, "数据表")


def parse(path: str, fmt: str = "xlsx") -> ParseResult:
    result = ParseResult(meta={"format": fmt})
    fmt = fmt.lower()

    try:
        if fmt == "csv":
            _parse_csv(path, result)
        elif fmt == "xlsx":
            _parse_xlsx(path, result)
        elif fmt == "xls":
            raise ParseError(
                "旧版 .xls 格式不支持直接解析，请在 Excel/WPS 中另存为 .xlsx 后重新上传。"
            )
        else:
            raise ParseError(f"不支持的表格格式：{fmt}")
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError(f"表格解析失败：{exc}") from exc

    if not result.blocks:
        raise ParseError("表格中没有解析出任何数据")
    return result
